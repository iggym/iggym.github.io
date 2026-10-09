"""Build the Word résumé from resume.json. Usage: python3 build_docx.py ../assets/resume/iggy-resume.docx"""
import json, sys
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches, RGBColor

import os
SRC = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'resume.json')
ACCENT = RGBColor(0x43, 0x38, 0xCA)
MUTED = RGBColor(0x5A, 0x5D, 0x66)

def rule(paragraph):
    """Thin bottom border under a section heading."""
    pPr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    for k, v in (('w:val', 'single'), ('w:sz', '6'), ('w:space', '1'), ('w:color', 'C9C9CF')):
        bottom.set(qn(k), v)
    borders.append(bottom)
    # Schema order: pBdr must come before spacing, ind and jc.
    anchor = next((c for c in pPr if c.tag in (qn('w:shd'), qn('w:tabs'), qn('w:spacing'), qn('w:ind'), qn('w:jc'), qn('w:rPr'))), None)
    if anchor is not None:
        anchor.addprevious(borders)
    else:
        pPr.append(borders)

def spacing(p, before=0, after=0):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)

def heading(doc, text):
    p = doc.add_paragraph()
    spacing(p, before=10, after=4)
    r = p.add_run(text.upper())
    r.bold = True; r.font.size = Pt(10); r.font.color.rgb = ACCENT
    rule(p)
    p.paragraph_format.keep_with_next = True
    return p

def main(out, src=SRC):
    d = json.load(open(src, encoding='utf-8'))
    compact = bool(d.get('compact'))
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(0.6 if compact else 0.75)
    sec.top_margin = sec.bottom_margin = Inches(0.5 if compact else 0.6)

    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'; normal.font.size = Pt(9.5 if compact else 10)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
    normal.paragraph_format.line_spacing = 1.08

    name = doc.add_paragraph(); spacing(name, after=0)
    r = name.add_run(d['name']); r.bold = True; r.font.size = Pt(22)
    h = doc.add_paragraph(); spacing(h, after=2)
    r = h.add_run(d['headline']); r.font.size = Pt(11); r.font.color.rgb = ACCENT
    c = d['contact']
    contact = doc.add_paragraph(); spacing(contact, after=2)
    contact.add_run(f"{d['location']}  ·  {c['phone']}  ·  {c['email']}  ·  {c['linkedin']}  ·  {c['web']}").font.size = Pt(9)

    heading(doc, 'Summary')
    for para in d['summary']:
        p = doc.add_paragraph(para); spacing(p, after=4)

    for label, text in d['focus']:
        p = doc.add_paragraph(style='List Bullet'); spacing(p, after=0)
        r = p.add_run(label + ': '); r.bold = True
        p.add_run(text)

    heading(doc, 'Experience')
    for role in d['roles']:
        p = doc.add_paragraph(); spacing(p, before=6, after=0)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(role['company']); r.bold = True; r.font.size = Pt(11)
        p = doc.add_paragraph(); spacing(p, after=0)
        p.paragraph_format.keep_with_next = True
        p.add_run(role['title']).italic = True
        meta = role['date'] + (f"  ·  {role['location']}" if role['location'] else '')
        p = doc.add_paragraph(); spacing(p, after=2)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(meta); r.font.size = Pt(9); r.font.color.rgb = MUTED
        for item in role['items']:
            if item['type'] == 'b':
                p = doc.add_paragraph(style='List Bullet'); spacing(p, after=1)
            else:
                p = doc.add_paragraph(); spacing(p, after=3)
            p.add_run(item['text'])

    if d.get('earlier'):
        heading(doc, 'Earlier experience')
        for e in d['earlier']:
            p = doc.add_paragraph(style='List Bullet'); spacing(p, after=0)
            p.add_run(e['company']).bold = True
            p.add_run((f" · {e['title']}" if e['title'] else '') + f" · {e['date']}")

    heading(doc, 'Education')
    for e in d['education']:
        p = doc.add_paragraph(); spacing(p, after=0)
        r = p.add_run(e['school']); r.bold = True
        p = doc.add_paragraph(); spacing(p, after=0)
        p.add_run(f"{e['degree']}  ·  {e['dates']}")

    heading(doc, 'Skills and languages')
    p = doc.add_paragraph(); spacing(p, after=0)
    p.add_run('Core skills: ').bold = True
    p.add_run(', '.join(d['top_skills']))
    p = doc.add_paragraph(); spacing(p, after=0)
    p.add_run('Languages: ').bold = True
    p.add_run(', '.join(d['languages']))

    doc.core_properties.title = f"{d['name']} résumé"
    doc.core_properties.author = d['name']
    doc.save(out)
    print('wrote', out)

if __name__ == '__main__':
    # usage: build_docx.py OUT.docx [SOURCE.json]
    main(sys.argv[1] if len(sys.argv) > 1 else 'iggy-resume.docx',
         sys.argv[2] if len(sys.argv) > 2 else SRC)
